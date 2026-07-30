from pathlib import Path

import openpyxl
from docx import Document
from pypdf import PdfReader

from app.models.schemas import Locator, TextSegment

SUPPORTED_SOURCE_TYPES = {"pdf", "docx", "txt", "xlsx"}


def load_pdf(file_path: Path) -> list[TextSegment]:
    reader = PdfReader(str(file_path))
    segments: list[TextSegment] = []
    order = 0
    for page_index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        segments.append(
            TextSegment(
                text=text,
                order=order,
                locator=Locator(kind="page", page_start=page_index, page_end=page_index),
            )
        )
        order += 1
    return segments


def load_docx(file_path: Path) -> list[TextSegment]:
    document = Document(str(file_path))
    segments: list[TextSegment] = []
    current_heading: str | None = None
    order = 0
    for para_index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style is not None and paragraph.style.name.lower().startswith("heading"):
            current_heading = text
            continue
        segments.append(
            TextSegment(
                text=text,
                order=order,
                locator=Locator(
                    kind="heading",
                    heading=current_heading,
                    para_start=para_index,
                    para_end=para_index,
                ),
            )
        )
        order += 1
    return segments


def load_excel(file_path: Path) -> list[TextSegment]:
    workbook = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
    segments: list[TextSegment] = []
    order = 0
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        try:
            header = [str(cell) if cell is not None else "" for cell in next(rows)]
        except StopIteration:
            continue
        for row_index, row in enumerate(rows, start=2):
            if row is None or all(cell is None for cell in row):
                continue
            cells = [str(cell) if cell is not None else "" for cell in row]
            # Row-wise serialization with the header repeated on every row, never flattened prose.
            text = " | ".join(f"{col}: {val}" for col, val in zip(header, cells))
            segments.append(
                TextSegment(
                    text=text,
                    order=order,
                    locator=Locator(kind="sheet_rows", sheet=sheet.title, row_start=row_index, row_end=row_index),
                )
            )
            order += 1
    return segments


def load_txt(file_path: Path) -> list[TextSegment]:
    raw = file_path.read_text(encoding="utf-8", errors="replace")
    segments: list[TextSegment] = []
    order = 0
    line_cursor = 1
    for block in raw.split("\n\n"):
        block_stripped = block.strip()
        line_count = block.count("\n") + 1
        if block_stripped:
            segments.append(
                TextSegment(
                    text=block_stripped,
                    order=order,
                    locator=Locator(
                        kind="text",
                        line_start=line_cursor,
                        line_end=line_cursor + line_count - 1,
                    ),
                )
            )
            order += 1
        line_cursor += line_count + 1  # account for the blank line consumed by the split
    return segments


_LOADERS = {
    "pdf": load_pdf,
    "docx": load_docx,
    "xlsx": load_excel,
    "txt": load_txt,
}


def extract_segments(file_path: Path, source_type: str) -> list[TextSegment]:
    if source_type not in _LOADERS:
        raise ValueError(f"unsupported source_type: {source_type!r}, expected one of {sorted(SUPPORTED_SOURCE_TYPES)}")
    segments = _LOADERS[source_type](file_path)
    if not segments:
        raise ValueError(f"no extractable text found in {file_path.name}")
    return segments
